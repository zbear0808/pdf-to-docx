"""Convert 'stats merged quiz unit 1.pdf' into high-fidelity editable DOCX and DocumentSpec AST."""

import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pdf_to_docx.ir_schema import (
    Alignment,
    DocumentSpec,
    HeadingBlock,
    ImageBlock,
    PageSpec,
    ParagraphBlock,
    TableBlock,
    TableCell,
    TextRun,
)


def set_cell_background(cell, hex_color: str):
    """Sets background shading of a table cell via OpenXML."""
    hex_clean = hex_color.lstrip("#")
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)
    shd_element = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_clean}" w:val="clear"/>')
    tcPr.append(shd_element)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Sets cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="000000", sz="4", val="single"):
    """Applies clean table borders to all outer and inner cell walls."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def make_row_cant_split(row):
    """Prevents a table row from breaking mid-cell across a page boundary (w:cantSplit)."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def make_row_header(row):
    """Repeats this table row at the top of each page if the table splits (w:tblHeader)."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


def set_box_container_borders(cell, sz="12"):
    """Applies a crisp outer frame border to the response box."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def build_ast() -> DocumentSpec:
    """Constructs the canonical DocumentSpec AST for the 7 pages."""
    doc_spec = DocumentSpec(
        title="AP Statistics Merged Quiz Unit 1",
        author="College Board / Educational Testing Service",
        theme_hex="#000000",
        default_font="Times New Roman",
        default_font_size_pt=10.5,
    )

    # ------------------ PAGE 1 ------------------
    p1 = PageSpec(page_number=1, width_pt=612.0, height_pt=792.0)
    p1.blocks.append(
        HeadingBlock(
            level=1,
            text="2011 AP\u00ae STATISTICS FREE-RESPONSE QUESTIONS",
            alignment=Alignment.CENTER,
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="2. The table below shows the political party registration by gender of all 500 registered voters in Franklin Township."
                )
            ]
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[TextRun(text="PARTY REGISTRATION\u2014FRANKLIN TOWNSHIP", bold=True)],
            alignment=Alignment.CENTER,
        )
    )
    p1.blocks.append(
        TableBlock(
            headers=["", "Party W", "Party X", "Party Y", "Total"],
            rows=[
                ["Female", "60", "120", "120", "300"],
                ["Male", "28", "124", "48", "200"],
                ["Total", "88", "244", "168", "500"],
            ],
            col_widths_pct=[22.0, 19.5, 19.5, 19.5, 19.5],
            cant_split=True,
            repeat_header=True,
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(a) ", bold=True),
                TextRun(
                    text="Given that a randomly selected registered voter is a male, what is the probability that he is registered for Party Y?"
                ),
            ]
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(b) ", bold=True),
                TextRun(
                    text="Among the registered voters of Franklin Township, are the events \u201cis a male\u201d and \u201cis registered for Party Y\u201d independent? Justify your answer based on probabilities calculated from the table above."
                ),
            ]
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(c) ", bold=True),
                TextRun(
                    text="One way to display the data in the table is to use a segmented bar graph. The following segmented bar graph, constructed from the data in the party registration\u2014Franklin Township table, shows party-registration distributions for males and females in Franklin Township."
                ),
            ]
        )
    )
    p1.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_1_img_1_xref4.png",
            caption="FRANKLIN TOWNSHIP Segmented Bar Graph",
            width_inches=5.8,
        )
    )
    p1.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="In Lawrence Township, the proportions of all registered voters for Parties W, X, and Y are the same as for Franklin Township, and party registration is independent of gender. Complete the graph below to show the distributions of party registration by gender in Lawrence Township."
                )
            ]
        )
    )
    p1.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_1_img_2_xref9.png",
            caption="LAWRENCE TOWNSHIP Response Graph",
            width_inches=5.8,
        )
    )
    doc_spec.pages.append(p1)

    # ------------------ PAGE 2 ------------------
    p2 = PageSpec(page_number=2, width_pt=594.6, height_pt=774.2)
    p2.blocks.append(
        HeadingBlock(
            level=2,
            text="Begin your response to QUESTION 1 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p2.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="STATISTICS\nSECTION II\nTotal Time\u20141 hour and 30 minutes\n6 Questions\n\nPart A\nQuestions 1-5\nSpend about 1 hour and 5 minutes on this part of the exam.", bold=True)
            ],
            alignment=Alignment.CENTER,
        )
    )
    p2.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="Directions: ", bold=True),
                TextRun(
                    text="Show all your work. Indicate clearly the methods you use, because you will be scored on the correctness of your methods as well as on the accuracy and completeness of your results and explanations."
                ),
            ]
        )
    )
    p2.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="1. The length of stay in a hospital after receiving a particular treatment is of interest to the patient, the hospital, and insurance providers. Of particular interest are unusually short or long lengths of stay. A random sample of 50 patients who received the treatment was selected, and the length of stay, in number of days, was recorded for each patient. The results are summarized in the following table and are shown in the dotplot."
                )
            ]
        )
    )
    p2.blocks.append(
        TableBlock(
            headers=["Length of stay (days)", "5", "6", "7", "8", "9", "12", "21"],
            rows=[["Number of patients", "4", "13", "14", "11", "6", "1", "1"]],
            cant_split=True,
            repeat_header=False,
        )
    )
    p2.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_2_img_1_xref139.png",
            caption="Length of Stay Dotplot",
            width_inches=3.2,
        )
    )
    p2.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(a) ", bold=True),
                TextRun(text="Determine the five-number summary of the distribution of length of stay."),
            ]
        )
    )
    doc_spec.pages.append(p2)

    # ------------------ PAGE 3 ------------------
    p3 = PageSpec(page_number=3, width_pt=594.6, height_pt=774.2)
    p3.blocks.append(
        HeadingBlock(
            level=2,
            text="Continue your response to QUESTION 1 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p3.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(b) ", bold=True),
                TextRun(
                    text="Consider two rules for identifying outliers, method A and method B. Let method A represent the 1.5 \u00d7 IQR rule, and let method B represent the 2 standard deviations rule."
                ),
            ]
        )
    )
    p3.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(i) ", bold=True),
                TextRun(
                    text="Using method A, determine any data points that are potential outliers in the distribution of length of stay. Justify your answer."
                ),
            ]
        )
    )
    p3.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(ii) ", bold=True),
                TextRun(
                    text="The mean length of stay for the sample is 7.42 days with a standard deviation of 2.37 days. Using method B, determine any data points that are potential outliers in the distribution of length of stay. Justify your answer."
                ),
            ]
        )
    )
    p3.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(c) ", bold=True),
                TextRun(
                    text="Explain why method A might identify more data points as potential outliers than method B for a distribution that is strongly skewed to the right."
                ),
            ]
        )
    )
    doc_spec.pages.append(p3)

    # ------------------ PAGE 4 ------------------
    p4 = PageSpec(page_number=4, width_pt=594.0, height_pt=774.0)
    p4.blocks.append(
        HeadingBlock(
            level=2,
            text="Begin your response to QUESTION 1 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p4.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="STATISTICS\nSECTION II\nTotal Time\u20141 hour and 30 minutes\n6 Questions\n\nPart A\nSuggested Time\u20141 hour and 5 minutes\n5 Questions",
                    bold=True,
                )
            ],
            alignment=Alignment.CENTER,
        )
    )
    p4.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="Directions: ", bold=True),
                TextRun(
                    text="Show all your work. Indicate clearly the methods you use, because you will be scored on the correctness of your methods as well as on the accuracy and completeness of your results and explanations."
                ),
            ]
        )
    )
    p4.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="1. As part of a study on the chemistry of Alaskan streams, researchers took water samples from many streams with temperatures colder than 8\u00b0C and from many streams with temperatures warmer than 8\u00b0C. For each sample, the researchers measured the dissolved oxygen concentration, in milligrams per liter (mg/l)."
                )
            ]
        )
    )
    p4.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_4_img_1_xref305.png",
            caption="Dissolved Oxygen Concentration Histogram",
            width_inches=5.6,
        )
    )
    p4.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(a) ", bold=True),
                TextRun(text="The researchers constructed the histogram shown for the dissolved oxygen concentration in streams from the sample with water temperatures "),
                TextRun(text="colder", underline=True),
                TextRun(text=" than 8\u00b0C. Based on the histogram, describe the distribution of dissolved oxygen concentration in streams with water temperatures "),
                TextRun(text="colder", underline=True),
                TextRun(text=" than 8\u00b0C."),
            ]
        )
    )
    doc_spec.pages.append(p4)

    # ------------------ PAGE 5 ------------------
    p5 = PageSpec(page_number=5, width_pt=594.0, height_pt=774.0)
    p5.blocks.append(
        HeadingBlock(
            level=2,
            text="Continue your response to QUESTION 1 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p5.blocks.append(
        TableBlock(
            headers=["Min", "Q1", "Median", "Q3", "Max", "Mean", "Std. Dev."],
            rows=[["2.10", "4.39", "5.43", "6.12", "13.45", "5.54", "1.64"]],
            cant_split=True,
            repeat_header=False,
        )
    )
    p5.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(b) ", bold=True),
                TextRun(text="The researchers computed the summary statistics shown in the table for the dissolved oxygen concentration in streams from the sample with water temperatures "),
                TextRun(text="warmer", underline=True),
                TextRun(text=" than 8\u00b0C. Use the summary statistics to construct a box plot for the dissolved oxygen concentration in streams with water temperatures "),
                TextRun(text="warmer", underline=True),
                TextRun(text=" than 8\u00b0C. Do not indicate outliers."),
            ]
        )
    )
    p5.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_5_img_1_xref319.png",
            caption="Box Plot Response Grid",
            width_inches=5.6,
        )
    )
    p5.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(c) ", bold=True),
                TextRun(text="The researchers believe that streams with higher dissolved oxygen concentration are generally healthier for wildlife. Which streams are generally healthier for wildlife, those with water temperature "),
                TextRun(text="colder", underline=True),
                TextRun(text=" than 8\u00b0C or those with water temperature "),
                TextRun(text="warmer", underline=True),
                TextRun(text=" than 8\u00b0C? Using characteristics of the distribution of dissolved oxygen concentration for temperatures "),
                TextRun(text="colder", underline=True),
                TextRun(text=" than 8\u00b0C and characteristics of the distribution of dissolved oxygen concentration for temperatures "),
                TextRun(text="warmer", underline=True),
                TextRun(text=" than 8\u00b0C, justify your answer."),
            ]
        )
    )
    doc_spec.pages.append(p5)

    # ------------------ PAGE 6 ------------------
    p6 = PageSpec(page_number=6, width_pt=612.0, height_pt=792.0)
    p6.blocks.append(
        HeadingBlock(
            level=2,
            text="Begin your response to QUESTION 2 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p6.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="2. A local elementary school decided to sell bottles printed with the school district\u2019s logo as a fund-raiser. The students in the elementary school were asked to sell bottles in three different sizes (small, medium, and large). The relative frequencies of the number of bottles sold for each size by the elementary school were 0.5 for small bottles, 0.3 for medium bottles, and 0.2 for large bottles."
                )
            ]
        )
    )
    p6.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="A local middle school also decided to sell bottles as a fund-raiser, using the same three sizes (small, medium, and large). The middle school students sold three times the number of bottles that the elementary school students sold. For the middle school students, the proportion of bottles sold was equal for all three sizes."
                )
            ]
        )
    )
    p6.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(a) ", bold=True),
                TextRun(
                    text="Complete the segmented bar graphs representing the relative frequencies of the number of bottles sold for each size by students at each school."
                ),
            ]
        )
    )
    p6.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_6_img_1_xref356.png",
            caption="Elementary & Middle School Segmented Bar Graphs Template",
            width_inches=4.5,
        )
    )
    p6.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(b) ", bold=True),
                TextRun(
                    text="An administrator at the elementary school concluded that the elementary school students sold more small bottles than the middle school students did. Is the elementary school administrator\u2019s conclusion correct? Explain your response."
                ),
            ]
        )
    )
    doc_spec.pages.append(p6)

    # ------------------ PAGE 7 ------------------
    p7 = PageSpec(page_number=7, width_pt=612.0, height_pt=792.0)
    p7.blocks.append(
        HeadingBlock(
            level=2,
            text="Continue your response to QUESTION 2 on this page.",
            alignment=Alignment.CENTER,
        )
    )
    p7.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="Two high schools are also selling the bottles and are competing to see which one sold more large bottles."
                )
            ]
        )
    )
    p7.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(c) ", bold=True),
                TextRun(
                    text="A mosaic plot for the distribution of the number of bottles sold by each of the high schools is shown here."
                ),
            ]
        )
    )
    p7.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(
                    text="Distribution of the Number of Bottles Sold by High School",
                    bold=True,
                )
            ],
            alignment=Alignment.CENTER,
        )
    )
    p7.blocks.append(
        ImageBlock(
            image_path="extracted_assets/page_7_img_1_xref369.png",
            caption="High School Bottles Sold Mosaic Plot",
            width_inches=4.6,
        )
    )
    p7.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(i) ", bold=True),
                TextRun(
                    text="Which of the two high schools sold a greater proportion of large bottles? Justify your answer."
                ),
            ]
        )
    )
    p7.blocks.append(
        ParagraphBlock(
            runs=[
                TextRun(text="(ii) ", bold=True),
                TextRun(
                    text="Which of the two high schools sold a greater number of large bottles? Justify your answer."
                ),
            ]
        )
    )
    doc_spec.pages.append(p7)

    return doc_spec


def compile_custom_docx(spec: DocumentSpec, output_path: Path):
    """Compiles the AP Statistics quiz into a high-fidelity Microsoft Word (.docx) document."""
    doc = Document()

    # Base typography
    normal_style = doc.styles["Normal"]
    font = normal_style.font
    font.name = "Times New Roman"
    font.size = Pt(10.0)
    font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    normal_style.paragraph_format.line_spacing = 1.10
    normal_style.paragraph_format.space_after = Pt(2)

    page_metadata = [
        {"year": "2011", "q": "2", "has_box": False, "page_label": "-7-", "response_lines": 0},
        {"year": "2021", "q": "1", "has_box": True, "page_label": "4", "response_lines": 4},
        {"year": "2021", "q": "1", "has_box": True, "page_label": "5", "response_lines": 4},
        {"year": "2023", "q": "1", "has_box": True, "page_label": "4", "response_lines": 4},
        {"year": "2023", "q": "1", "has_box": True, "page_label": "5", "response_lines": 3},
        {"year": "2024", "q": "2", "has_box": True, "page_label": "5", "response_lines": 4},
        {"year": "2024", "q": "2", "has_box": True, "page_label": "6", "response_lines": 4},
    ]

    for p_idx, page in enumerate(spec.pages):
        meta = page_metadata[p_idx]

        if p_idx == 0:
            section = doc.sections[0]
        else:
            section = doc.add_section()

        # Geometry
        section.top_margin = Inches(0.40)
        section.bottom_margin = Inches(0.40)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.50)
        section.header_distance = Inches(0.20)
        section.footer_distance = Inches(0.20)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        hp.paragraph_format.space_after = Pt(0)
        if meta["year"] != "2011":
            hrun = hp.add_run(f"AP\u00ae Statistics {meta['year']} Free-Response Questions")
            hrun.font.name = "Times New Roman"
            hrun.font.size = Pt(9.0)
            hrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)

        if meta["year"] == "2021":
            f_warn = fp.add_run("Use a pencil or pen with black or dark blue ink only. Do NOT write your name. Do NOT write outside the box.\n")
            f_warn.font.size = Pt(7.5)
            f_warn.font.name = "Times New Roman"

        f_copy = fp.add_run(f"\u00a9 {meta['year']} College Board.   Visit College Board on the web: collegeboard.org.\n")
        f_copy.font.size = Pt(7.5)
        f_copy.font.name = "Times New Roman"
        f_copy.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        f_num = fp.add_run(meta["page_label"])
        f_num.font.size = Pt(9.0)
        f_num.font.name = "Times New Roman"

        # Now build content for this page
        if meta["has_box"]:
            # Wrap in single-cell container table with crisp border
            container_tbl = doc.add_table(rows=1, cols=1)
            container_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            container_tbl.autofit = False
            container_tbl.columns[0].width = Inches(7.45)

            cell = container_tbl.cell(0, 0)
            set_box_container_borders(cell, sz="12")  # 1.5pt solid black border
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

            first_p = cell.paragraphs[0]
            first_p_used = False

            def add_p_to_container(alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=Pt(0), space_after=Pt(2)):
                nonlocal first_p_used
                if not first_p_used:
                    p = first_p
                    first_p_used = True
                else:
                    p = cell.add_paragraph()
                p.alignment = alignment
                p.paragraph_format.space_before = space_before
                p.paragraph_format.space_after = space_after
                p.paragraph_format.line_spacing = 1.08
                return p

            for block in page.blocks:
                b_type = getattr(block, "type", "")
                if b_type == "heading":
                    p = add_p_to_container(
                        alignment=WD_ALIGN_PARAGRAPH.CENTER,
                        space_before=Pt(2),
                        space_after=Pt(4),
                    )
                    r = p.add_run(block.text)
                    r.bold = True
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10.5)
                elif b_type == "paragraph":
                    align = WD_ALIGN_PARAGRAPH.LEFT
                    if block.alignment == Alignment.CENTER:
                        align = WD_ALIGN_PARAGRAPH.CENTER
                    p = add_p_to_container(alignment=align, space_before=Pt(2), space_after=Pt(3))
                    for r_spec in block.runs:
                        r = p.add_run(r_spec.text)
                        r.bold = r_spec.bold
                        r.italic = r_spec.italic
                        r.underline = r_spec.underline
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10.0)
                elif b_type == "table":
                    all_rows = [block.headers] + block.rows
                    sub_table = cell.add_table(rows=len(all_rows), cols=len(all_rows[0]))
                    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(sub_table, sz="5")
                    for r_idx, r_data in enumerate(all_rows):
                        row = sub_table.rows[r_idx]
                        make_row_cant_split(row)
                        for c_idx, val in enumerate(r_data):
                            sub_cell = row.cells[c_idx]
                            set_cell_margins(sub_cell, top=50, bottom=50, left=70, right=70)
                            cp = sub_cell.paragraphs[0]
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cp.paragraph_format.space_before = Pt(0)
                            cp.paragraph_format.space_after = Pt(0)
                            r = cp.add_run(val)
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(9.5)
                            if r_idx == 0:
                                r.bold = True
                elif b_type == "image":
                    p = add_p_to_container(
                        alignment=WD_ALIGN_PARAGRAPH.CENTER,
                        space_before=Pt(4),
                        space_after=Pt(4),
                    )
                    img_w = Inches(block.width_inches) if block.width_inches else Inches(5.2)
                    p.add_run().add_picture(block.image_path, width=img_w)

            # Add "GO ON TO THE NEXT PAGE." banner inside the bottom of the box
            bot_p = cell.add_paragraph()
            bot_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            bot_p.paragraph_format.space_before = Pt(8)
            bot_p.paragraph_format.space_after = Pt(2)
            r_bot = bot_p.add_run("GO ON TO THE NEXT PAGE.")
            r_bot.bold = True
            r_bot.font.name = "Times New Roman"
            r_bot.font.size = Pt(9.5)

        else:
            # Page 1 (No box frame, direct document flow)
            for block in page.blocks:
                b_type = getattr(block, "type", "")
                if b_type == "heading":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(4)
                    r = p.add_run(block.text)
                    r.bold = True
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11.5)
                elif b_type == "paragraph":
                    align = WD_ALIGN_PARAGRAPH.LEFT
                    if block.alignment == Alignment.CENTER:
                        align = WD_ALIGN_PARAGRAPH.CENTER
                    p = doc.add_paragraph()
                    p.alignment = align
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.08
                    for r_spec in block.runs:
                        r = p.add_run(r_spec.text)
                        r.bold = r_spec.bold
                        r.italic = r_spec.italic
                        r.underline = r_spec.underline
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10.0)
                elif b_type == "table":
                    all_rows = [block.headers] + block.rows
                    table = doc.add_table(rows=len(all_rows), cols=len(all_rows[0]))
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, sz="5")
                    for r_idx, r_data in enumerate(all_rows):
                        row = table.rows[r_idx]
                        make_row_cant_split(row)
                        for c_idx, val in enumerate(r_data):
                            t_cell = row.cells[c_idx]
                            set_cell_margins(t_cell, top=45, bottom=45, left=70, right=70)
                            cp = t_cell.paragraphs[0]
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cp.paragraph_format.space_before = Pt(0)
                            cp.paragraph_format.space_after = Pt(0)
                            r = cp.add_run(val)
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(9.5)
                            if r_idx == 0 or c_idx == 0:
                                r.bold = True
                elif b_type == "image":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    img_w = Inches(block.width_inches) if block.width_inches else Inches(5.6)
                    p.add_run().add_picture(block.image_path, width=img_w)

            # Bottom page 1 banner
            bot_p = doc.add_paragraph()
            bot_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            bot_p.paragraph_format.space_before = Pt(6)
            bot_p.paragraph_format.space_after = Pt(0)
            r_bot = bot_p.add_run("GO ON TO THE NEXT PAGE.")
            r_bot.bold = True
            r_bot.font.name = "Times New Roman"
            r_bot.font.size = Pt(9.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Successfully compiled DOCX to: {output_path} ({output_path.stat().st_size} bytes)")


def main():
    ast_spec = build_ast()

    # Save AST JSON
    ast_file = Path("tests/stats_merged_quiz_unit_1_ast.json")
    ast_file.write_text(ast_spec.model_dump_json(indent=2), encoding="utf-8")
    print(f"Saved canonical AST to: {ast_file}")

    # Compile DOCX
    out_docx = Path("tests/stats merged quiz unit 1.docx")
    compile_custom_docx(ast_spec, out_docx)

    # Also copy to test_output for reference in its own independent directory
    test_dir = Path("test_output/stats_merged_quiz")
    test_dir.mkdir(parents=True, exist_ok=True)
    out_copy = test_dir / "stats merged quiz unit 1.docx"
    out_copy.write_bytes(out_docx.read_bytes())
    ast_copy = test_dir / "stats_merged_quiz_unit_1_ast.json"
    ast_copy.write_text(ast_spec.model_dump_json(indent=2), encoding="utf-8")
    print(f"Copied to test_output folder: {test_dir}")


if __name__ == "__main__":
    main()
