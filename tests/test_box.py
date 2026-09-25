"""Test python-docx capabilities for exam box and nested tables."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.5)
sec.left_margin = Inches(0.6)
sec.right_margin = Inches(0.6)

# Test 1x1 container table
table = doc.add_table(rows=1, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(7.1)

cell = table.cell(0, 0)
# Set border on cell
tcPr = cell._tc.get_or_add_tcPr()
borders = parse_xml(
    f'<w:tcBorders {nsdecls("w")}>\n'
    f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
    f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
    f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
    f'  <w:right w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
    f'</w:tcBorders>'
)
tcPr.append(borders)

p0 = cell.paragraphs[0]
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
r0 = p0.add_run("Begin your response to QUESTION 1 on this page.")
r0.bold = True

p1 = cell.add_paragraph()
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p1.add_run("STATISTICS\nSECTION II")
r1.bold = True

# Add a nested table inside cell
nested = cell.add_table(rows=2, cols=3)
nested.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in nested.rows:
    for c in r.cells:
        c.text = "data"

test_path = Path("test_output/test_box.docx")
doc.save(str(test_path))
print(f"Saved {test_path} ({test_path.stat().st_size} bytes)")
