"""Deterministic DOCX Compiler.

Transforms a DocumentSpec AST into a production-grade Word document (.docx).
Enforces OpenXML rules for repeating table headers, atomic table rows,
and twip-precise column layouts.
"""

from __future__ import annotations
import logging
from pathlib import Path
from typing import Optional
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from .ir_schema import (
    Alignment, ChartBlock, DocumentBlock, DocumentSpec,
    HeadingBlock, ImageBlock, ListType, PageBreakBlock,
    PageSpec, ParagraphBlock, TableBlock, TextRun,
)
from .chart_generator import render_chart_image


def hex_to_rgb(hex_str: str) -> RGBColor:
    hex_clean = hex_str.lstrip("#")
    if len(hex_clean) == 3:
        hex_clean = "".join([c * 2 for c in hex_clean])
    return RGBColor(int(hex_clean[0:2], 16), int(hex_clean[2:4], 16), int(hex_clean[4:6], 16))


def set_cell_background(cell, hex_color: str):
    hex_clean = hex_color.lstrip("#")
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)
    shd_element = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_clean}" w:val="clear"/>')
    tcPr.append(shd_element)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def sanitize_xml(text: str) -> str:
    """Removes XML control characters that crash python-docx or Word."""
    if not isinstance(text, str):
        return str(text) if text is not None else ""
    return "".join(
        c for c in text
        if c in ("\t", "\n", "\r")
        or (0x20 <= ord(c) <= 0xD7FF)
        or (0xE000 <= ord(c) <= 0xFFFD)
        or (0x10000 <= ord(c) <= 0x10FFFF)
    )


def set_table_borders(table, color="444444", sz="4", val="single"):
    """Applies crisp table borders to all outer and inner cell walls."""
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


class DocxCompiler:
    """Compiles a canonical DocumentSpec AST into a formatted Word (.docx) file."""

    def __init__(self, spec: DocumentSpec):
        self.spec = spec
        self.doc = Document()
        self._configure_styles()

    def _configure_styles(self):
        normal_style = self.doc.styles["Normal"]
        font = normal_style.font
        font.name = self.spec.default_font
        font.size = Pt(self.spec.default_font_size_pt)
        font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        normal_style.paragraph_format.line_spacing = 1.15
        normal_style.paragraph_format.space_after = Pt(4)

    def compile(self, output_docx_path: str | Path, assets_dir: Optional[str | Path] = None) -> Path:
        out_path = Path(output_docx_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        assets_path = Path(assets_dir) if assets_dir else out_path.parent / "assets"
        assets_path.mkdir(parents=True, exist_ok=True)

        for page_idx, page in enumerate(self.spec.pages):
            if page_idx > 0:
                self.doc.add_page_break()
            for block in page.blocks:
                self._render_block(block, assets_path)

        try:
            self.doc.save(str(out_path))
            return out_path
        except PermissionError:
            fallback = out_path.parent / f"{out_path.stem}_clean{out_path.suffix}"
            logger.warning(
                f"File '{out_path}' is locked by another process (e.g. Word). "
                f"Saving to fallback path: '{fallback}'"
            )
            self.doc.save(str(fallback))
            return fallback

    def _render_block(self, block: DocumentBlock, assets_path: Path):
        b_type = getattr(block, "type", "")
        if b_type == "heading":
            self._render_heading(block)
        elif b_type == "paragraph":
            self._render_paragraph(block)
        elif b_type == "table":
            self._render_table(block)
        elif b_type == "chart":
            self._render_chart(block, assets_path)
        elif b_type == "image":
            self._render_image(block)
        elif b_type == "page_break":
            self.doc.add_page_break()

    def _render_heading(self, h: HeadingBlock):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(max(6, 18 - (h.level * 2)))
        p.paragraph_format.space_after = Pt(4)
        if h.alignment == Alignment.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif h.alignment == Alignment.RIGHT:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if h.runs:
            for run in h.runs:
                self._apply_run(p.add_run(sanitize_xml(run.text)), run, is_heading=True, level=h.level)
        else:
            run = p.add_run(sanitize_xml(h.text))
            run.bold = True
            run.font.size = Pt(max(12, 20 - (h.level * 2)))
            if self.spec.theme_hex and self.spec.theme_hex.upper() not in ("#000000", "#111111", "#222222", "#1F4E79"):
                run.font.color.rgb = hex_to_rgb(self.spec.theme_hex)
            else:
                run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    def _render_paragraph(self, p_block: ParagraphBlock):
        p = self.doc.add_paragraph()
        if p_block.list_type == ListType.BULLET:
            p.style = "List Bullet"
        elif p_block.list_type == ListType.NUMBERED:
            p.style = "List Number"
        if p_block.is_callout:
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        if p_block.alignment == Alignment.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p_block.alignment == Alignment.RIGHT:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif p_block.alignment == Alignment.JUSTIFY:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r_spec in p_block.runs:
            self._apply_run(p.add_run(sanitize_xml(r_spec.text)), r_spec)

    def _apply_run(self, docx_run, r_spec: TextRun, is_heading: bool = False, level: int = 1):
        docx_run.bold = r_spec.bold
        docx_run.italic = r_spec.italic
        docx_run.underline = r_spec.underline
        docx_run.font.strike = r_spec.strike
        if r_spec.font_size_pt:
            docx_run.font.size = Pt(r_spec.font_size_pt)
        elif is_heading:
            docx_run.font.size = Pt(max(12, 20 - (level * 2)))
        if r_spec.color_hex:
            docx_run.font.color.rgb = hex_to_rgb(r_spec.color_hex)
        elif is_heading:
            if self.spec.theme_hex and self.spec.theme_hex.upper() not in ("#000000", "#111111", "#222222", "#1F4E79"):
                docx_run.font.color.rgb = hex_to_rgb(self.spec.theme_hex)
            else:
                docx_run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        if r_spec.font_name:
            docx_run.font.name = r_spec.font_name

    def _render_table(self, tbl: TableBlock):
        all_rows = []
        if tbl.headers:
            all_rows.append([sanitize_xml(h) for h in tbl.headers])

        sanitized_rows = [[sanitize_xml(c) for c in r] for r in tbl.rows]
        # Deduplicate if rows[0] is identical to headers
        if tbl.headers and sanitized_rows:
            h_norm = [str(h).strip().lower() for h in tbl.headers]
            r0_norm = [str(c).strip().lower() for c in sanitized_rows[0]]
            if h_norm == r0_norm:
                sanitized_rows = sanitized_rows[1:]
        all_rows.extend(sanitized_rows)

        if not all_rows:
            return

        num_rows = len(all_rows)
        num_cols = max(len(r) for r in all_rows)
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="444444", sz="4")

        is_themed = bool(self.spec.theme_hex and self.spec.theme_hex.upper() not in ("#000000", "#111111", "#222222", "#1F4E79"))

        for row_idx, row_data in enumerate(all_rows):
            doc_row = table.rows[row_idx]
            make_row_cant_split(doc_row)
            is_header = row_idx == 0 and bool(tbl.headers)
            if is_header:
                make_row_header(doc_row)
            for col_idx in range(num_cols):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = doc_row.cells[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                p = cell.paragraphs[0]
                p.text = ""
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if col_idx == 0 and not is_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif is_header or cell_text.replace(".", "").isdigit():
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(cell_text)
                if is_header:
                    run.bold = True
                    if is_themed:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        set_cell_background(cell, self.spec.theme_hex)
                    else:
                        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
                        set_cell_background(cell, "#F2F4F7")
                    run.font.size = Pt(9.5)
                else:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "#FAFAFA")

        if tbl.col_widths_pct and len(tbl.col_widths_pct) == num_cols:
            total_width = Inches(6.5)
            for col_idx, pct in enumerate(tbl.col_widths_pct):
                col_w = total_width * (pct / 100.0)
                for row in table.rows:
                    row.cells[col_idx].width = col_w

    def _render_chart(self, chart: ChartBlock, assets_path: Path):
        chart_img_path = chart.generated_image_path
        if not chart_img_path or not Path(chart_img_path).exists():
            chart_filename = f"chart_{id(chart)}.png"
            chart_img_path = str(assets_path / chart_filename)
            render_chart_image(chart, chart_img_path)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        shape = self.doc.add_picture(str(chart_img_path), width=Inches(min(6.0, chart.width_inches)))
        chart_alt = sanitize_xml((chart.title or "").strip() or "Data chart")
        try:
            docPr = shape._inline.find(qn("wp:docPr"))
            if docPr is not None:
                docPr.set("descr", chart_alt)
                docPr.set("title", chart_alt)
        except Exception:
            pass

    def _render_image(self, img: ImageBlock):
        img_path = Path(img.image_path) if img.image_path else None
        if not img_path or not img_path.exists():
            return
        p = self.doc.add_paragraph()
        if img.alignment == Alignment.LEFT:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif img.alignment == Alignment.RIGHT:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)

        # Inspect true image pixel dimensions to strictly preserve aspect ratio
        aspect_ratio = None
        try:
            with Image.open(img_path) as im:
                px_w, px_h = im.size
                if px_h > 0:
                    aspect_ratio = px_w / px_h
        except Exception as e:
            logger.warning(f"Could not read dimensions from {img_path}: {e}")

        # Standard page printable constraints (inches)
        MAX_PAGE_WIDTH = 6.5
        MAX_PAGE_HEIGHT = 8.5

        # Determine target display width in inches (clamped to page content width)
        target_w = float(img.width_inches) if (img.width_inches and img.width_inches > 0) else 5.5
        target_w = min(target_w, MAX_PAGE_WIDTH)

        # If aspect ratio is known, ensure height does not overflow page content height
        if aspect_ratio:
            projected_h = target_w / aspect_ratio
            if projected_h > MAX_PAGE_HEIGHT:
                target_w = MAX_PAGE_HEIGHT * aspect_ratio

        # Pass ONLY width to add_picture so python-docx computes proportional height,
        # guaranteeing 100% exact aspect ratio preservation with zero distortion
        shape = self.doc.add_picture(str(img_path), width=Inches(target_w))

        # Set accessibility alt text on DrawingML docPr
        img_alt = sanitize_xml((img.caption or "").strip() or "Document figure")
        try:
            docPr = shape._inline.find(qn("wp:docPr"))
            if docPr is not None:
                docPr.set("descr", img_alt)
                docPr.set("title", img_alt)
        except Exception:
            pass

        if img.caption:
            cap_p = self.doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(6)
            run = cap_p.add_run(sanitize_xml(img.caption))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
