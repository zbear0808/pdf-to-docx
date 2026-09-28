"""Deterministic DOCX Compiler.

Transforms a DocumentSpec AST into a production-grade Word document (.docx).
Enforces OpenXML rules for repeating table headers (w:tblHeader), atomic table rows
(w:cantSplit), twip-precise column layouts, and robust typography.
"""

from __future__ import annotations
from pathlib import Path
from typing import Any, Dict, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from .ir_schema import (
    Alignment,
    BoxBlock,
    ChartBlock,
    CodeBlock,
    DocumentBlock,
    DocumentSpec,
    EquationBlock,
    HeadingBlock,
    ImageBlock,
    ListType,
    PageBreakBlock,
    PageSpec,
    ParagraphBlock,
    QuestionBlock,
    TableBlock,
    TableCell,
    TextRun,
)
from .chart_generator import render_chart_image


def hex_to_rgb(hex_str: str) -> RGBColor:
    """Converts hex color string (#RRGGBB) to docx RGBColor."""
    hex_clean = hex_str.lstrip("#")
    if len(hex_clean) == 3:
        hex_clean = "".join([c * 2 for c in hex_clean])
    r = int(hex_clean[0:2], 16)
    g = int(hex_clean[2:4], 16)
    b = int(hex_clean[4:6], 16)
    return RGBColor(r, g, b)


def set_cell_background(cell, hex_color: str):
    """Sets background shading of a table cell via OpenXML."""
    hex_clean = hex_color.lstrip("#")
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if present
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)
    shd_element = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_clean}" w:val="clear"/>')
    tcPr.append(shd_element)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in twips (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def make_row_cant_split(row):
    """Prevents a table row from breaking mid-cell across a page boundary (w:cantSplit)."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def make_row_header(row):
    """Repeats this table row at the top of each page if the table splits (w:tblHeader)."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


class DocxCompiler:
    """Compiles a canonical DocumentSpec AST into a formatted Word (.docx) file."""

    def __init__(self, spec: DocumentSpec):
        self.spec = spec
        self.doc = Document()
        self._configure_styles()

    def _configure_styles(self):
        """Sets up default typographic styles with corporate theme colors."""
        normal_style = self.doc.styles["Normal"]
        font = normal_style.font
        font.name = self.spec.default_font
        font.size = Pt(self.spec.default_font_size_pt)
        font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        normal_style.paragraph_format.line_spacing = 1.05
        normal_style.paragraph_format.space_after = Pt(2)

    def _apply_page_geometry(self, section, page: PageSpec):
        """Applies exact PDF page dimensions and margins to Word section."""
        if page.width_pt and page.height_pt:
            section.page_width = Pt(page.width_pt)
            section.page_height = Pt(page.height_pt)

        if page.orientation == "landscape":
            from docx.enum.section import WD_ORIENTATION
            section.orientation = WD_ORIENTATION.LANDSCAPE
            if page.width_pt and page.height_pt and page.width_pt < page.height_pt:
                section.page_width = Pt(page.height_pt)
                section.page_height = Pt(page.width_pt)

        top_m = page.margin_top_pt if page.margin_top_pt is not None else (self.spec.margins_in * 72 if self.spec.margins_in else 36.0)
        btm_m = page.margin_bottom_pt if page.margin_bottom_pt is not None else (self.spec.margins_in * 72 if self.spec.margins_in else 36.0)
        left_m = page.margin_left_pt if page.margin_left_pt is not None else (self.spec.margins_in * 72 if self.spec.margins_in else 40.0)
        right_m = page.margin_right_pt if page.margin_right_pt is not None else (self.spec.margins_in * 72 if self.spec.margins_in else 40.0)

        section.top_margin = Pt(max(18.0, float(top_m)))
        section.bottom_margin = Pt(max(18.0, float(btm_m)))
        section.left_margin = Pt(max(24.0, float(left_m)))
        section.right_margin = Pt(max(24.0, float(right_m)))

    def compile(self, output_docx_path: str | Path, assets_dir: Optional[str | Path] = None) -> Path:
        """Executes full compilation and writes output file."""
        out_path = Path(output_docx_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        assets_path = Path(assets_dir) if assets_dir else out_path.parent / "assets"
        assets_path.mkdir(parents=True, exist_ok=True)

        if self.spec.pages:
            self._apply_page_geometry(self.doc.sections[0], self.spec.pages[0])

        for page_idx, page in enumerate(self.spec.pages):
            if page_idx > 0:
                prev_page = self.spec.pages[page_idx - 1]
                need_new_section = (
                    page.orientation != prev_page.orientation
                    or abs((page.width_pt or 612) - (prev_page.width_pt or 612)) > 10
                    or abs((page.height_pt or 792) - (prev_page.height_pt or 792)) > 10
                )
                if need_new_section:
                    sec = self.doc.add_section()
                    self._apply_page_geometry(sec, page)
                else:
                    self.doc.add_page_break()

            for block in page.blocks:
                self._render_block(block, assets_path)

        self.doc.save(str(out_path))
        return out_path

    def _render_block(self, block: DocumentBlock, assets_path: Path):
        b_type = getattr(block, "type", "")

        if b_type == "heading":
            self._render_heading(block)  # type: ignore
        elif b_type == "paragraph":
            self._render_paragraph(block)  # type: ignore
        elif b_type == "table":
            self._render_table(block)  # type: ignore
        elif b_type == "chart":
            self._render_chart(block, assets_path)  # type: ignore
        elif b_type == "image":
            self._render_image(block)  # type: ignore
        elif b_type == "equation":
            self._render_equation(block)  # type: ignore
        elif b_type == "box":
            self._render_box(block)  # type: ignore
        elif b_type == "question":
            self._render_question(block)  # type: ignore
        elif b_type == "code":
            self._render_code(block)  # type: ignore
        elif b_type == "page_break":
            self.doc.add_page_break()

    def _render_heading(self, h: HeadingBlock):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(max(4, 14 - (h.level * 2)))
        p.paragraph_format.space_after = Pt(3)

        if h.alignment == Alignment.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif h.alignment == Alignment.RIGHT:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        if h.runs:
            for run in h.runs:
                self._apply_run(p.add_run(run.text), run, is_heading=True, level=h.level)
        else:
            run = p.add_run(h.text)
            run.bold = True
            run.font.size = Pt(max(12, 18 - (h.level * 2)))
            run.font.color.rgb = hex_to_rgb(self.spec.theme_hex)

    def _render_paragraph(self, p_block: ParagraphBlock):
        p = self.doc.add_paragraph()

        # Handle list styles
        if p_block.list_type == ListType.BULLET:
            p.style = "List Bullet"
        elif p_block.list_type == ListType.NUMBERED:
            p.style = "List Number"

        # Handle callout styling
        if p_block.is_callout:
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif p_block.space_before_pt is not None:
            p.paragraph_format.space_before = Pt(p_block.space_before_pt)

        if p_block.space_after_pt is not None:
            p.paragraph_format.space_after = Pt(p_block.space_after_pt)
        else:
            p.paragraph_format.space_after = Pt(2)

        if p_block.line_spacing is not None:
            p.paragraph_format.line_spacing = p_block.line_spacing

        # Handle alignment
        if p_block.alignment == Alignment.CENTER:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p_block.alignment == Alignment.RIGHT:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif p_block.alignment == Alignment.JUSTIFY:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for r_spec in p_block.runs:
            self._apply_run(p.add_run(r_spec.text), r_spec)

    def _render_equation(self, eq: EquationBlock):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        code = eq.latex_code.strip()
        for delim in ("$$", "\\[", "\\]"):
            code = code.replace(delim, "").strip()
        run = p.add_run(code)
        run.italic = True
        run.font.name = "Cambria Math"

    def _render_box(self, box: BoxBlock):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        border_color = box.color_hex.lstrip("#") if box.color_hex else "444444"
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:left w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:right w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)
        if box.title:
            p0 = cell.paragraphs[0]
            p0.paragraph_format.space_after = Pt(2)
            r = p0.add_run(box.title)
            r.bold = True
        first_p = not bool(box.title)
        for p_spec in box.content:
            p = cell.paragraphs[0] if first_p else cell.add_paragraph()
            first_p = False
            p.paragraph_format.space_after = Pt(2)
            for r_spec in p_spec.runs:
                self._apply_run(p.add_run(r_spec.text), r_spec)
        if box.empty_for_response:
            table.rows[0].height = Pt(box.height_pt or 100.0)

    def _render_question(self, q: QuestionBlock):
        q_prefix = ""
        if q.number is not None:
            q_prefix = f"{q.number}. "
        if q.part:
            q_prefix += f"({q.part}) "
        for i, p_spec in enumerate(q.prompt):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if i == 0 and q_prefix:
                r_pre = p.add_run(q_prefix)
                r_pre.bold = True
            for r_spec in p_spec.runs:
                self._apply_run(p.add_run(r_spec.text), r_spec)
        if q.choices:
            for c in q.choices:
                p_ch = self.doc.add_paragraph()
                p_ch.paragraph_format.left_indent = Inches(0.25)
                p_ch.paragraph_format.space_after = Pt(2)
                r_lbl = p_ch.add_run(f"{c.label} ")
                r_lbl.bold = True
                if c.runs:
                    for r_spec in c.runs:
                        self._apply_run(p_ch.add_run(r_spec.text), r_spec)
                else:
                    p_ch.add_run(c.text)
        if q.response_box_height_pt and q.response_box_height_pt > 0:
            self._render_box(BoxBlock(height_pt=q.response_box_height_pt, empty_for_response=True))

    def _render_code(self, code_block: CodeBlock):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "#F4F5F7")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_block.code)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)

    def _apply_run(self, docx_run, r_spec: TextRun, is_heading: bool = False, level: int = 1):
        docx_run.bold = r_spec.bold
        docx_run.italic = r_spec.italic
        docx_run.underline = r_spec.underline
        docx_run.font.strike = r_spec.strike

        if r_spec.font_size_pt:
            docx_run.font.size = Pt(r_spec.font_size_pt)
        elif is_heading:
            docx_run.font.size = Pt(max(12, 22 - (level * 2)))

        if r_spec.color_hex:
            docx_run.font.color.rgb = hex_to_rgb(r_spec.color_hex)
        elif is_heading:
            docx_run.font.color.rgb = hex_to_rgb(self.spec.theme_hex)

        if r_spec.font_name:
            docx_run.font.name = r_spec.font_name

    def _render_table(self, tbl: TableBlock):
        all_rows = []
        if tbl.headers:
            all_rows.append(tbl.headers)
        all_rows.extend(tbl.rows)

        if not all_rows:
            return

        num_rows = len(all_rows)
        num_cols = max(len(r) for r in all_rows)

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if tbl.style_name:
            try:
                table.style = tbl.style_name
            except Exception:
                table.style = "Table Grid"

        # Apply header formatting & OOXML rules
        for row_idx, row_data in enumerate(all_rows):
            doc_row = table.rows[row_idx]

            # Enforce atomic rows (never split across pages)
            if tbl.cant_split:
                make_row_cant_split(doc_row)

            # Enforce header row repetition
            is_header = row_idx == 0 and bool(tbl.headers)
            if is_header and tbl.repeat_header:
                make_row_header(doc_row)

            for col_idx in range(num_cols):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = doc_row.cells[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

                # Clear default paragraph
                p = cell.paragraphs[0]
                p.text = ""
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0

                run = p.add_run(cell_text)

                if is_header:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(10)
                    set_cell_background(cell, self.spec.theme_hex)
                else:
                    run.font.size = Pt(9.5)
                    # Alternate row shading for readability
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "#F7F9FB")

        # Explicit column widths if available
        if tbl.col_widths_pct and len(tbl.col_widths_pct) == num_cols:
            total_width = Inches(6.5)  # Standard printable width on 8.5" page with 1" margins
            for col_idx, pct in enumerate(tbl.col_widths_pct):
                col_w = total_width * (pct / 100.0)
                for row in table.rows:
                    row.cells[col_idx].width = col_w

    def _render_chart(self, chart: ChartBlock, assets_path: Path):
        chart_img_path = chart.generated_image_path
        if not chart_img_path or not Path(chart_img_path).exists():
            chart_filename = f"chart_{id(chart)}.png"
            chart_img_path = str(assets_path / chart_filename)
            render_chart_image(chart, chart_img_path)

        if not Path(chart_img_path).exists():
            return

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        self.doc.add_picture(str(chart_img_path), width=Inches(min(6.0, chart.width_inches)))

    def _render_image(self, img: ImageBlock):
        if not img.image_path or not Path(img.image_path).exists():
            return

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)

        w = Inches(img.width_inches) if img.width_inches else Inches(5.5)
        h = Inches(img.height_inches) if img.height_inches else None
        if h:
            self.doc.add_picture(img.image_path, width=w, height=h)
        else:
            self.doc.add_picture(img.image_path, width=w)

        if img.caption:
            cap_p = self.doc.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(6)
            run = cap_p.add_run(img.caption)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
